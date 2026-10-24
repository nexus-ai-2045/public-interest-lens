#!/usr/bin/env python3
"""国策の数値目標（data/acquired/fiscal-more/policy-targets.csv）の達成状況を測る系列を、官公庁の一次資料から取り込み、目標と結び付ける（*.go.jp）。

段（取得の段はそれぞれ単独で回せる。毎朝の自動更新は --monthly で段ごとに月1回。どれも年1回〜四半期の公表）
- sna：内閣府 経済社会総合研究所「四半期別GDP速報」の年度・暦年の表（支出側の名目・実質の額、前年度比）。
- pb：内閣府「中長期の経済財政に関する試算」の計数表（国・地方の基礎的財政収支。骨太の方針の PB 目標の物差し）。実績の年度だけ。
- wage：厚生労働省「地域別最低賃金改定状況」の全国加重平均額（年度）。
- visitors：日本政府観光局（JNTO）「年別 訪日外客数」（暦年）。
- spending：観光庁「訪日外国人消費動向調査（インバウンド消費動向調査）」の暦年の概要（旅行消費額、暦年）。
- lodging：観光庁「宿泊旅行統計調査」推移表の都道府県別 外国人延べ宿泊者数（年計）。地方部（三大都市圏8都府県以外）と東北6県を足し上げる。
- fdi：財務省「本邦対外資産負債残高の推移」の負債の直接投資（暦年末。対日直接投資残高の目標の物差し＝資産負債原則）。
- exports：農林水産省「農林水産物・食品の輸出額」（暦年。年ごとの概要 PDF の1頁目）。
- rd：総務省「科学技術研究調査」の研究費の総額（年度。統計ダッシュボードの API）。
- ghg：国立環境研究所 温室効果ガスインベントリオフィス「日本の温室効果ガス排出量データ」（年度。総排出量と、吸収量を差し引いた排出・吸収量）。
- link：手元の段。上の系列（と DB の消費者物価指数）から計算する系列を作り、目標ごとに「定義の合う系列」を結び付けて達成状況を出す。
  定義の合わない目標は無理に結ばず、理由を target-links.csv に残す（結び付けの規則は RULES）。
- all：上の段を順に。

出力は data/acquired/target-series/ の CSV（Git に入れる）。系列は段ごとの <段>.csv（落ちた段は前回のファイルを残す）、計算した系列は derived.csv、
目標との結び付けは target-links.csv。原本は .local/target-series/raw/ に置き、manifest.csv に URL・sha256・取得日を残し、
.local/external-materials/target-series/ にも外部資料として登録する（AGENTS.md）。DB への取り込みは scripts/build_database.py（load_series・load_target_progress）。

使い方：python3 scripts/fetch_target_series.py all [--monthly] [--dry-run] [--offline]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import sys
import time
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlsplit

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.fetch_fiscal_more import (USER_AGENT, links, nfkc, number, read_csv, record_fetch, squash,  # noqa: E402
                                       write_csv)
from scripts import fetch_fiscal_more  # noqa: E402

OUT = ROOT / "data" / "acquired" / "target-series"
RAW = ROOT / ".local" / "target-series" / "raw"
TARGETS = ROOT / "data" / "acquired" / "fiscal-more" / "policy-targets.csv"
CPI = ROOT / "data" / "acquired" / "domestic-stats" / "cpi_items_annual_2020base.csv"
LAST_FETCH = OUT / "last-fetch.txt"
MONTHLY_DAYS = 28

SNA_TOP = "https://www.esri.cao.go.jp/jp/sna/sokuhou/sokuhou_top.html"
SHISAN_PAGE = "https://www5.cao.go.jp/keizai2/keizai-syakai/shisan.html"
WAGE_PAGE = "https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/koyou_roudou/roudoukijun/minimumichiran/"
JNTO_PAGE = "https://www.jnto.go.jp/statistics/data/visitors-statistics/"
SPENDING_PAGE = "https://www.mlit.go.jp/kankocho/tokei_hakusyo/gaikokujinshohidoko.html"
LODGING_PAGE = "https://www.mlit.go.jp/kankocho/tokei_hakusyo/shukuhakutokei.html"
IIP_XLS = "https://www.mof.go.jp/policy/international_policy/reference/iip/6IIP.xls"
EXPORTS_PAGE = "https://www.maff.go.jp/j/shokusan/export/e_info/zisseki.html"
RD_API = "https://dashboard.e-stat.go.jp/api/1.0/Json/getData?Lang=JP&IndicatorCode=1102010000000010000&MetaGetFlg=N"
RD_PAGE = "https://www.stat.go.jp/data/kagaku/index.html"
GHG_PAGE = "https://www.nies.go.jp/gio/archive/ghgdata/index.html"

SERIES_HEADER = ["系列ID", "系列", "年", "値", "単位", "年の別", "出典URL", "備考"]
LINK_HEADER = ["決定ID", "指標", "目標値", "単位", "目標年", "年の別", "系列ID", "達成の値", "達成の年", "判定", "理由"]
MANIFEST_HEADER = ["URL", "sha256", "bytes", "取得日", "保存先"]
# 三大都市圏（観光庁の報道発表の定義）。地方部はこれ以外の道県
METRO = ("埼玉", "千葉", "東京", "神奈川", "愛知", "京都", "大阪", "兵庫")
TOHOKU = ("青森", "岩手", "宮城", "秋田", "山形", "福島")
CHO, OKU, MAN = 1e12, 1e8, 1e4


# ---------- 共通 ----------

class Fetcher:
    """原本を .local/target-series/raw/ に保存し、manifest.csv に記録する。--offline は保存済みの原本だけを読む。"""

    def __init__(self, offline: bool = False, wait: float = 0.5):
        self.offline = offline
        self.wait = wait
        self.manifest = {r["URL"]: r for r in read_csv(OUT / "manifest.csv")}

    def path(self, url: str) -> Path:
        parts = urlsplit(url)
        name = parts.path.lstrip("/") or "index.html"
        if name.endswith("/"):
            name += "index.html"
        if parts.query:  # API（統計ダッシュボード）は問い合わせごとに別のファイルにする
            name += "-" + hashlib.sha256(parts.query.encode()).hexdigest()[:12] + ".json"
        return RAW / parts.netloc / name

    def get(self, url: str) -> bytes:
        target = self.path(url)
        if self.offline:
            return target.read_bytes()
        request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(request, timeout=90) as response:
            body = response.read()
        time.sleep(self.wait)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(body)
        self.manifest[url] = {"URL": url, "sha256": hashlib.sha256(body).hexdigest(), "bytes": str(len(body)),
                              "取得日": date.today().isoformat(), "保存先": str(target.relative_to(ROOT))}
        register(target, url)
        return body

    def file(self, url: str) -> Path:
        self.get(url)
        return self.path(url)

    def text(self, url: str) -> str:
        body = self.get(url)
        try:
            return body.decode("utf-8")
        except UnicodeDecodeError:
            return body.decode("cp932", "replace")

    def save_manifest(self) -> None:
        write_csv(OUT / "manifest.csv", MANIFEST_HEADER, [self.manifest[k] for k in sorted(self.manifest)])


def register(path: Path, url: str) -> None:
    """データの原本を .local/external-materials/target-series/ に登録する（AGENTS.md）。"""
    from scripts.register_external_material import register_once

    register_once(input_path=path, source_url=url, task_id="target-series", title=f"数値目標の達成状況の原本（{path.name}）",
                  purpose="国策の数値目標の達成状況を測る系列をデータベースに入れる", storage_root=ROOT / ".local" / "external-materials",
                  observed_at=datetime.now(timezone.utc).isoformat(timespec="seconds"))


def due(phase: str, today: date | None = None) -> bool:
    return fetch_fiscal_more.due(today=today, stamp=LAST_FETCH, days=MONTHLY_DAYS, phase=phase)


def pdf_text(path: Path, last_page: int | None = None) -> str:
    args = ["pdftotext", "-layout"] + (["-l", str(last_page)] if last_page else []) + [str(path), "-"]
    return subprocess.run(args, capture_output=True, text=True, check=True).stdout


def xlsx_rows(path: Path, sheet: str | None = None) -> list[list]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)  # read_only でないとグラフの定義で落ちる表がある（GHG）
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def xls_rows(path: Path) -> list[list]:
    import xlrd

    sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
    return [sh.row_values(i) for i in range(sh.nrows)]


def yen(text: str) -> float | None:
    """「1兆7,005億円」「9,860億円」→ 円。"""
    m = re.fullmatch(r"(?:(\d+)兆)?([\d,]+)億円", squash(text))
    if not m:
        return None
    return (int(m.group(1) or 0) * 10000 + int(m.group(2).replace(",", ""))) * OKU


def row(sid: str, label: str, year: int, value: float, unit: str, kind: str, url: str, note: str = "") -> dict:
    value = round(value) if unit in ("円", "人", "人泊", "t-CO2") else round(value, 6)  # 額・人数は桁の端数（浮動小数の誤差）を残さない
    return {"系列ID": sid, "系列": label, "年": year, "値": value, "単位": unit, "年の別": kind, "出典URL": url, "備考": note}


def _not_empty(phase: str, rows: list) -> None:
    if not rows:
        raise RuntimeError(f"{phase}：行が取れない（前回の取得結果を残す）")


# ---------- 国民経済計算（内閣府） ----------

SNA_FILES = {  # ファイルの頭: (系列ID, 系列名, 単位, 年の別, 備考)
    "gaku-mfy": ("sna-qe:名目GDP:年度", "国内総生産（支出側、名目）", "円", "年度", "10億円を円に"),
    "gaku-mcy": ("sna-qe:名目GDP:暦年", "国内総生産（支出側、名目）", "円", "暦年", "10億円を円に"),
    "gaku-jfy": ("sna-qe:実質GDP:年度", "国内総生産（支出側、実質、連鎖方式）", "円", "年度", "10億円を円に"),
    "ritu-mfy": ("sna-qe:名目GDP成長率:年度", "国内総生産（支出側、名目）の前年度比", "%", "年度", ""),
    "ritu-jfy": ("sna-qe:実質GDP成長率:年度", "国内総生産（支出側、実質）の前年度比", "%", "年度", ""),
}


def parse_sna_csv(text: str, sid: str, label: str, unit: str, kind: str, url: str, note: str) -> list[dict]:
    """1列目が「1994/4-3.」（年度）か「1994/1-12.」（暦年）、2列目が国内総生産（支出側）。"""
    import csv
    import io

    out = []
    for cells in csv.reader(io.StringIO(text)):
        if len(cells) < 2:
            continue
        m = re.match(r"(\d{4})/(4-3|1-12)\.", nfkc(cells[0]))
        v = number(cells[1])
        if not m or v is None:
            continue
        out.append(row(sid, label, int(m.group(1)), v * 1e9 if unit == "円" else v, unit, kind, url, note))
    return out


def fetch_sna(fetcher: Fetcher) -> dict:
    top = fetcher.text(SNA_TOP)
    menu = next((u for u, _ in links(top, SNA_TOP) if u.endswith("gdemenuja.html")), None)
    m = re.search(r"/qe(\d{3})_(\d)/", menu or "")
    if not m:
        raise RuntimeError("sna：最新の速報のページが見つからない")
    code, base = m.group(1) + m.group(2), menu.rsplit("/", 1)[0] + "/tables/"
    rows = []
    for head, (sid, label, unit, kind, note) in SNA_FILES.items():
        url = f"{base}{head}{code}.csv"
        body = fetcher.get(url)
        rows += parse_sna_csv(body.decode("cp932", "replace"), sid, label, unit, kind, url, "。".join(x for x in (note, f"四半期別GDP速報 {code} の年次値") if x))
    _not_empty("sna", rows)
    write_csv(OUT / "sna.csv", SERIES_HEADER, rows)
    return {"行": len(rows), "速報": code}


# ---------- 国・地方の基礎的財政収支（内閣府 中長期試算） ----------

def parse_pb(rows: list[list], url: str, edition_year: int, title: str) -> list[dict]:
    """「年度」の行と「基礎的財政収支」の行（とその次の対名目GDP比の行）。実績は試算の版の年の2年前の年度まで（以降は推計と見通し）。"""
    years_row = next((r for r in rows if squash(r[0] if r else "") == "年度"), None)
    pb_i = next((i for i, r in enumerate(rows) if r and squash(r[0]) == "基礎的財政収支"), None)
    if years_row is None or pb_i is None:
        raise RuntimeError("pb：年度か基礎的財政収支の行が見つからない")
    last_actual = edition_year - 2
    out = []
    for j, y in enumerate(years_row):
        if not isinstance(y, (int, float)) or y > last_actual:
            continue
        pb, ratio = number(rows[pb_i][j]), number(rows[pb_i + 1][j])
        note = f"{title}。実績（決算）の年度だけ"
        if pb is not None:
            out.append(row("pb-cao:国・地方:兆円", "国・地方の基礎的財政収支（中長期試算）", int(y), pb * CHO, "円", "年度", url, note))
        if ratio is not None:
            out.append(row("pb-cao:国・地方:対GDP比", "国・地方の基礎的財政収支の対名目GDP比（中長期試算）", int(y), ratio, "%", "年度", url, note))
    return out


def fetch_pb(fetcher: Fetcher) -> dict:
    page = fetcher.text(SHISAN_PAGE)
    found = links(page, SHISAN_PAGE)
    xlsx = next((u for u, _ in found if u.endswith("keisuuhyo.xlsx")), None)
    hontai = next(((u, t) for u, t in found if re.search(r"(\d{2})(\d{2})hontai\.pdf$", u)), None)
    if not xlsx or not hontai:
        raise RuntimeError("pb：計数表か本体の PDF が見つからない")
    edition_year = 2000 + int(re.search(r"(\d{2})\d{2}hontai\.pdf$", hontai[0]).group(1))  # 「2607hontai.pdf」＝2026年7月の版
    path = fetcher.file(xlsx)
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(n for n in wb.sheetnames if "財政の詳細計数表" in n)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    title = next((nfkc(r[0]) for r in rows if r and isinstance(r[0], str) and "国・地方の財政の姿" in r[0]), "国・地方の財政の姿")
    out = parse_pb(rows, xlsx, edition_year, f"{re.split(r'[（(]PDF', nfkc(hontai[1]))[0]} {title}")
    _not_empty("pb", out)
    write_csv(OUT / "pb.csv", SERIES_HEADER, out)
    return {"行": len(out), "版": edition_year}


# ---------- 最低賃金（厚生労働省） ----------

def parse_wage(rows: list[list], url: str) -> list[dict]:
    """1行目に年度（「平成14年度」…）、「全国加重平均額」の行の各年度の「改定額」の列（年度の列）。"""
    head = rows[0]
    avg = next(r for r in rows if r and squash(r[0]) == "全国加重平均額")
    out = []
    for j, h in enumerate(head):
        y = fetch_fiscal_more.wareki_year(nfkc(h or ""))
        v = number(avg[j]) if j < len(avg) else None
        if y and v:
            out.append(row("minwage:全国加重平均", "地域別最低賃金の全国加重平均額（改定後）", y, round(v), "円", "年度", url,
                           "その年度の改定（10〜12月に発効）の額。最新の年度は答申・決定の額で、まだ発効していない県を含むことがある。小数のセルは四捨五入（公表の額）"))
    return out


def fetch_wage(fetcher: Fetcher) -> dict:
    page = fetcher.text(WAGE_PAGE)
    url = next((u for u, t in links(page, WAGE_PAGE) if u.endswith(".xlsx") and "地域別最低賃金改定状況" in t), None)
    if not url:
        raise RuntimeError("wage：改定状況の表が見つからない")
    out = parse_wage(xlsx_rows(fetcher.file(url)), url)
    _not_empty("wage", out)
    write_csv(OUT / "wage.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 訪日外客数（JNTO） ----------

def parse_visitors(text: str, url: str) -> list[dict]:
    out = []
    for line in text.splitlines():
        m = re.match(r"\s*(\d{4})\s+\S.*?年\s+([\d,]+)\s", nfkc(line))
        if m:
            out.append(row("jnto:訪日外客数", "訪日外客数", int(m.group(1)), float(m.group(2).replace(",", "")), "人", "暦年", url, "確定値"))
    return out


def fetch_visitors(fetcher: Fetcher) -> dict:
    page = fetcher.text(JNTO_PAGE)
    url = next((u for u, t in links(page, JNTO_PAGE) if u.endswith(".pdf") and "年別" in t and "訪日外客数" in t), None)
    if not url:
        raise RuntimeError("visitors：年別の推移の PDF が見つからない")
    out = parse_visitors(pdf_text(fetcher.file(url)), url)
    _not_empty("visitors", out)
    write_csv(OUT / "visitors.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 訪日外国人旅行消費額（観光庁） ----------

def spending_pages(html: str, base: str) -> list[tuple[int, str]]:
    """「暦年」「年間試算結果」の PDF と、その前の「2025年」の見出しの年。"""
    out = []
    for m in re.finditer(r'<a\b[^>]*href="([^"]+\.pdf)"[^>]*>(.*?)</a>', html, re.S | re.I):
        text = re.sub(r"<[^>]+>|\s+", "", m.group(2))
        if text not in ("暦年", "年間試算結果"):
            continue
        years = re.findall(r"(20\d{2})年", re.sub(r"<[^>]+>", " ", html[max(0, m.start() - 3000):m.start()]))
        if years:
            out.append((int(years[-1]), urljoin(base, m.group(1))))
    return out


def parse_spending(text: str, year: int, url: str) -> dict | None:
    t = re.sub(r"\s+", "", nfkc(text))
    if not re.search(rf"{year}年(?:暦年|の訪日外国人旅行消費額)", t[:300]):  # 「2024年比」などではなく見出しの年
        return None
    m = re.search(r"旅行消費額(?:\([^)]*\))?(?:について[^。]*?)?は((?:\d+兆)?[\d,]+億円)", t)
    value = yen(m.group(1)) if m else None
    if value is None:
        return None
    note = "試算値（調査の一部を中止）" if "試算" in t[:400] else "確報"
    if year >= 2018:
        note += "。2018年からクルーズ客を含む"
    return row("kanko:訪日外国人旅行消費額", "訪日外国人旅行消費額", year, value, "円", "暦年", url, note)


def fetch_spending(fetcher: Fetcher) -> dict:
    page = fetcher.text(SPENDING_PAGE)
    out = []
    for year, url in sorted(set(spending_pages(page, SPENDING_PAGE))):
        r = parse_spending(pdf_text(fetcher.file(url), last_page=1), year, url)
        if r:
            out.append(r)
    _not_empty("spending", out)
    write_csv(OUT / "spending.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 外国人延べ宿泊者数（観光庁） ----------

def parse_lodging(rows: list[list], url: str) -> list[dict]:
    """推移表の「旧3-1」：3行目に年（平成23年…）、「全国」と「01北海道」…「47沖縄県」の行（人泊）。"""
    years_row = next(r for r in rows if sum(1 for c in r if fetch_fiscal_more.wareki_year(nfkc(c or ""))) >= 3)
    years = {j: fetch_fiscal_more.wareki_year(nfkc(c or "")) for j, c in enumerate(years_row)}
    prefs = {}
    for r in rows:
        name = squash(r[0] if r else "")
        if name == "全国" or re.match(r"\d{2}\D", name):
            prefs[re.sub(r"^\d{2}", "", name)] = r
    if "全国" not in prefs or len(prefs) < 48:
        raise RuntimeError("lodging：全国と47都道府県の行が揃わない")
    out = []
    note = "宿泊旅行統計調査（全施設、確定値）。2026年1月分から標本の層化基準が変わった"
    for j, y in years.items():
        if not y:
            continue
        vals = {k: number(v[j]) for k, v in prefs.items()}
        if any(v is None for v in vals.values()):
            continue
        local = vals["全国"] - sum(v for k, v in vals.items() if k != "全国" and k.startswith(METRO))
        tohoku = sum(v for k, v in vals.items() if k.startswith(TOHOKU))
        out += [row("kanko:外国人延べ宿泊者数:全国", "外国人延べ宿泊者数（全国）", y, vals["全国"], "人泊", "暦年", url, note),
                row("kanko:外国人延べ宿泊者数:地方部", "外国人延べ宿泊者数（地方部＝三大都市圏8都府県以外）", y, local, "人泊", "暦年", url,
                    note + "。全国から埼玉・千葉・東京・神奈川・愛知・京都・大阪・兵庫を引いた"),
                row("kanko:外国人延べ宿泊者数:東北6県", "外国人延べ宿泊者数（東北6県）", y, tohoku, "人泊", "暦年", url,
                    note + "。青森・岩手・宮城・秋田・山形・福島の足し上げ")]
    return out


def fetch_lodging(fetcher: Fetcher) -> dict:
    page = fetcher.text(LODGING_PAGE)
    url = next((u for u, t in links(page, LODGING_PAGE) if u.endswith(".xlsx") and squash(t).startswith("推移表")), None)
    if not url:
        raise RuntimeError("lodging：推移表が見つからない")
    out = parse_lodging(xlsx_rows(fetcher.file(url), "旧3-1"), url)
    _not_empty("lodging", out)
    write_csv(OUT / "lodging.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 対日直接投資残高（財務省） ----------

def parse_iip(rows: list[list], url: str) -> list[dict]:
    """「負債残高」の下の「直接投資」の列（10億円）。行は「End of 1996」…（暦年末）。"""
    head_i = next(i for i, r in enumerate(rows) if "負債残高" in [nfkc(c) for c in r])
    liab = [nfkc(c) for c in rows[head_i]].index("負債残高")
    col = next(j for j in range(liab, len(rows[head_i + 2])) if nfkc(rows[head_i + 2][j]) == "直接投資")
    out = []
    for r in rows[head_i + 3:]:
        m = re.match(r"End of (\d{4})", nfkc(r[1] if len(r) > 1 else ""))
        v = number(r[col]) if m and col < len(r) else None
        if m and v is not None:
            out.append(row("mof-iip:負債:直接投資", "対日直接投資残高（本邦対外資産負債残高の負債の直接投資、資産負債原則）", int(m.group(1)), v * 1e9, "円",
                           "暦年", url, "暦年末。BPM6 準拠（2013年末以前は BPM5 からの組み替え）。毎年改訂される"))
    return out


def fetch_fdi(fetcher: Fetcher) -> dict:
    out = parse_iip(xls_rows(fetcher.file(IIP_XLS)), IIP_XLS)
    _not_empty("fdi", out)
    write_csv(OUT / "fdi.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 農林水産物・食品の輸出額（農林水産省） ----------

def export_pages(html: str, base: str) -> list[tuple[int, str]]:
    """年ごとの「農林水産物・食品の輸出額」「輸出実績（確定値）」の PDF（国・地域別・品目別・少額貨物・月の分は除く）。"""
    out = []
    for u, t in links(html, base):
        t = nfkc(t)
        m = re.search(r"[(（](\d{4})年[)）]", t)
        if not u.endswith(".pdf") or not m or re.search(r"国・地域別|品目別|少額|訪日|分割|\d+月", t):
            continue
        if ("輸出額" in t or "確定値" in t) and int(m.group(1)) >= 2018:  # 2017年以前は1頁目の書き方が違い読めないので取らない
            out.append((int(m.group(1)), u))
    return out


def parse_export(text: str, year: int, url: str) -> dict | None:
    """1頁目の年計（「2025年1-12月の輸出額は…1兆7,005億円」「1-12月累計 9,860億円」「2019年の輸出額は9,121億円」）。"""
    t = re.sub(r"\s+", "", nfkc(text))
    amount = r"((?:\d+兆)?[\d,]+億円)"
    m = (re.search(rf"{year}年(?:1-12月)?の輸出(?:額|実績)は[、,]?(?:対前年比[^%]*%の)?{amount}", t)
         or re.search(rf"1-12月累計{amount}", t))
    value = yen(m.group(1)) if m else None
    if value is None:
        return None
    note = "2020年から少額貨物（1品目20万円以下）を含む" if year >= 2020 else "少額貨物を含まない（当時の定義）"
    return row("maff:農林水産物・食品の輸出額", "農林水産物・食品の輸出額", year, value, "円", "暦年", url, note)


def fetch_exports(fetcher: Fetcher) -> dict:
    page = fetcher.text(EXPORTS_PAGE)
    out = []
    for year, url in sorted(dict(sorted(export_pages(page, EXPORTS_PAGE), reverse=True)).items()):
        try:
            r = parse_export(pdf_text(fetcher.file(url), last_page=1), year, url)
        except subprocess.CalledProcessError:
            r = None
        if r:
            out.append(r)
    _not_empty("exports", out)
    write_csv(OUT / "exports.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 研究費（総務省 科学技術研究調査） ----------

def parse_rd(payload: dict, url: str) -> list[dict]:
    objs = payload["GET_STATS"]["STATISTICAL_DATA"]["DATA_INF"]["DATA_OBJ"]
    out = []
    for o in objs:
        v = o["VALUE"]
        m = re.match(r"(\d{4})FY", v.get("@time", ""))
        val = number(v.get("$"))
        if m and val is not None:
            out.append(row("stat-kagaku:研究費総額", "研究費の総額（科学技術研究調査）", int(m.group(1)), val * 1e6, "円", "年度", url,
                           "百万円を円に。統計ダッシュボードの API（総務省統計局）"))
    return out


def fetch_rd(fetcher: Fetcher) -> dict:
    out = parse_rd(json.loads(fetcher.get(RD_API)), RD_API)
    _not_empty("rd", out)
    write_csv(OUT / "rd.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 温室効果ガス（国立環境研究所） ----------

def parse_ghg(rows: list[list], url: str) -> list[dict]:
    """「1.Summary」：年の行（1990…）の下の「計」（総排出量）と、「■排出・吸収量」の下の「排出・吸収量（計）」（百万t CO2換算）。"""
    def label(r):
        return "".join(squash(c) for c in r[:26] if isinstance(c, str))

    years_i = next(i for i, r in enumerate(rows) if 1990 in r)
    years = {j: int(c) for j, c in enumerate(rows[years_i]) if isinstance(c, (int, float)) and 1990 <= c <= 2100}
    total = next(r for r in rows[years_i:] if label(r).startswith("計"))
    net = next(r for r in rows if label(r).startswith("排出・吸収量(計)"))
    out = []
    for j, y in years.items():
        if number(total[j]) is not None:
            out.append(row("gio:温室効果ガス:総排出量", "温室効果ガスの総排出量", y, number(total[j]) * 1e6, "t-CO2", "年度", url, "確報値。吸収量を差し引かない"))
        if number(net[j]) is not None:
            out.append(row("gio:温室効果ガス:排出・吸収量", "温室効果ガスの排出・吸収量（森林等の吸収源対策による吸収量を差し引いた量）", y, number(net[j]) * 1e6,
                           "t-CO2", "年度", url, "確報値。吸収量の計上は2014年度から"))
    return out


def fetch_ghg(fetcher: Fetcher) -> dict:
    page = fetcher.text(GHG_PAGE)
    cands = [(int(m.group(1)), u) for u, _ in links(page, GHG_PAGE) if (m := re.search(r"L5-7gas_(\d{4})_gioweb", u))]
    if not cands:
        raise RuntimeError("ghg：排出量データの Excel が見つからない")
    url = max(cands)[1]
    out = parse_ghg(xlsx_rows(fetcher.file(url), "1.Summary"), url)
    _not_empty("ghg", out)
    write_csv(OUT / "ghg.csv", SERIES_HEADER, out)
    return {"行": len(out)}


# ---------- 計算する系列 ----------

def series_table() -> dict[str, dict[int, dict]]:
    table: dict[str, dict[int, dict]] = {}
    for path in sorted(OUT.glob("*.csv")):
        if path.stem in FETCH_PHASES:
            for r in read_csv(path):
                table.setdefault(r["系列ID"], {})[int(r["年"])] = r
    return table


def derive(table: dict[str, dict[int, dict]]) -> list[dict]:
    """取り込んだ系列から計算する系列（前年比・比率・2013年度比）。"""
    out = []

    def yoy(sid: str, new: str, label: str, kind: str, note: str):
        s = table.get(sid, {})
        for y in sorted(s):
            if y - 1 in s and float(s[y - 1]["値"]):
                out.append(row(new, label, y, (float(s[y]["値"]) / float(s[y - 1]["値"]) - 1) * 100, "%", kind, s[y]["出典URL"], note))

    yoy("minwage:全国加重平均", "minwage:全国加重平均:引上げ率", "地域別最低賃金の全国加重平均額の引上げ率", "年度", "全国加重平均額（円）の前年度比から計算")
    cpi = {}
    for r in read_csv(CPI):
        if r["品目"] in ("総合", "生鮮食品を除く総合"):
            cpi.setdefault(r["品目"], {})[int(r["年"])] = r
    for item, s in cpi.items():
        for y in sorted(s):
            if y - 1 in s:
                out.append(row(f"cpi-yoy:2020基準:{item}", f"消費者物価指数（{item}）の前年比", y, (float(s[y]["値"]) / float(s[y - 1]["値"]) - 1) * 100,
                               "%", "暦年", s[y]["出典URL"], "2020年基準の年平均指数（小数1桁）から計算。公表の前年比と0.1ポイントずれることがある"))
    rd, gdp = table.get("stat-kagaku:研究費総額", {}), table.get("sna-qe:名目GDP:年度", {})
    for y in sorted(set(rd) & set(gdp)):
        out.append(row("stat-kagaku:研究費総額:対GDP比", "研究費の総額の対名目GDP比", y, float(rd[y]["値"]) / float(gdp[y]["値"]) * 100, "%", "年度",
                       rd[y]["出典URL"], "研究費の総額÷名目GDP（年度、支出側、最新の速報の2020年基準）。GDP の基準改定の前の年度は総務省の公表の比と0.1ポイントほどずれる"))
    for sid in ("gio:温室効果ガス:総排出量", "gio:温室効果ガス:排出・吸収量"):
        s, base = table.get(sid, {}), table.get("gio:温室効果ガス:総排出量", {}).get(2013)
        if not base:
            continue
        for y in sorted(s):
            out.append(row(f"{sid}:2013年度比の削減率", f"{s[y]['系列']}の2013年度比の削減率", y, (1 - float(s[y]["値"]) / float(base["値"])) * 100, "%", "年度",
                           s[y]["出典URL"], "分母は2013年度の総排出量（環境省の公表と同じ）"))
    return out


# ---------- 目標との結び付け ----------

# (指標の正規表現, 単位の頭, 系列ID か None, 測り方, 理由・注記)。上から順に最初に当たった規則を使う。系列ID が None の規則は「結ばない」。
# 測り方：level＝その年の値、surplus＝0以上で黒字化、avg10＝2013〜目標年度の平均成長率（額の幾何平均）、
# rate＝毎年の率（成長率・物価上昇率・引上げ率）。目標年があればその年の率、無ければ決定の翌年から最新の年までの平均（1年だけの反動で「到達」としない）
GDP_FY, GDP_CY = "sna-qe:名目GDP:年度", "sna-qe:名目GDP:暦年"
RULES = [
    (r"基礎的財政収支|プライマリーバランス|PB", "黒字化", "pb-cao:国・地方:兆円", "surplus",
     "達成の値は兆円。内閣府 中長期試算の国・地方 PB（復旧・復興対策などの経費と財源を除くベース。除く範囲は試算の版で変わる）"),
    (r"名目GDP成長率\(10年平均\)", "%", "sna-qe:名目GDP:年度", "avg10", "名目GDP（支出側）の10年間の年平均成長率"),
    (r"実質GDP成長率\(10年平均\)", "%", "sna-qe:実質GDP:年度", "avg10", "実質GDP（支出側、連鎖）の10年間の年平均成長率"),
    (r"名目(?:GDP|経済)成長率", "%", "sna-qe:名目GDP成長率:年度", "rate", "名目GDP（支出側）の前年度比"),
    (r"実質(?:GDP|経済)成長率", "%", "sna-qe:実質GDP成長率:年度", "rate", "実質GDP（支出側）の前年度比"),
    (r"1人当たり名目国民総所得", "", None, "", "起点の年と1人当たりの人口の定義が本文の範囲に無い"),
    (r"物価", "%", "cpi-yoy:2020基準:総合", "rate",
     "消費者物価指数（総合）の前年比（日本銀行の目標の文言「消費者物価の前年比上昇率」）。生鮮食品を除く総合は cpi-yoy:2020基準:生鮮食品を除く総合。持続的・安定的かは別に見る"),
    (r"名目GDP|名目経済規模", "兆円", "?gdp", "level",
     "名目GDP（支出側、2020年基準。目標を決めた当時の2015年基準より水準が高い）。目標の年の別に合わせ、年なら暦年、年度か年の無いものは年度の系列"),
    (r"最低賃金.*引上げ率", "%", "minwage:全国加重平均:引上げ率", "rate", "全国加重平均額の前年度比"),
    (r"最低賃金", "円", "minwage:全国加重平均", "level", "地域別最低賃金の全国加重平均額（年度の改定後）"),
    (r"訪日(?:外国人旅行者数|客数)", "万人", "jnto:訪日外客数", "level", "JNTO の訪日外客数（暦年、確定値）"),
    (r"訪日(?:外国人)?(?:旅行者)?の?(?:旅行)?消費額|訪日客の消費額", "兆円", "kanko:訪日外国人旅行消費額", "level",
     "観光庁の訪日外国人旅行消費額（暦年）。2018年からクルーズ客を含み、2020〜2022年は試算値"),
    (r"地方部の外国人延べ宿泊者数", "万人泊", "kanko:外国人延べ宿泊者数:地方部", "level", "三大都市圏8都府県以外の外国人延べ宿泊者数（暦年）"),
    (r"東北6県の外国人宿泊者数", "万人泊", "kanko:外国人延べ宿泊者数:東北6県", "level", "東北6県の外国人延べ宿泊者数（暦年）"),
    (r"対(?:日|内)直接投資残高", "兆円", "mof-iip:負債:直接投資", "level", "本邦対外資産負債残高の負債の直接投資（資産負債原則。政府の目標の物差し）"),
    (r"農林水産物・食品の輸出額|^輸出額\(5兆円目標\)", "兆円", "maff:農林水産物・食品の輸出額", "level",
     "農林水産物・食品の輸出額（暦年）。2020年から少額貨物を含む（それ以前の目標の定義とは少し違う）"),
    (r"官民.*研究開発投資.*対GDP比", "%", "stat-kagaku:研究費総額:対GDP比", "level",
     "科学技術研究調査の研究費の総額÷名目GDP（年度）。目標は第5期科学技術基本計画の「2020年度まで」で、本文の「年」は年度で測る"),
    (r"政府.*研究開発投資.*対GDP比|政府研究開発投資\(対GDP比\)", "", None, "",
     "政府研究開発投資は内閣府 CSTI の科学技術関係予算（予算ベース、地方分を含む）で測るが、年度別の対GDP比の実績が表として公表されていない"),
    (r"企業から大学", "", None, "", "倍率の起点の年と、企業から大学・国立研究開発法人への研究費の系列を取り込んでいない"),
    (r"研究開発投資|研究開発投資総額", "", None, "", "5年間などの累計の目標。科学技術関係予算の年度別の表（予算ベース）を取り込んでいない"),
    (r"温室効果ガス.*2013年度比", "%", "gio:温室効果ガス:排出・吸収量:2013年度比の削減率", "level",
     "吸収量を差し引いた排出・吸収量の2013年度の総排出量に対する削減率（環境省が46%目標と並べて示す測り方）"),
    (r"温室効果ガス", "", None, "", "削減率の基準の年が本文の範囲に無い"),
]
DEFAULT_REASON = "定義の合う系列を取り込んでいない（所管の統計がこの取り込みの対象外）"


def base_unit(unit: str) -> str:
    return re.sub(r"以上|超|程度|未満|（[^）]*）", "", unit)


def scale(unit: str) -> float:
    u = base_unit(unit)
    return {"兆円": CHO, "億円": OKU, "万人": MAN, "万人泊": MAN}.get(u, 1.0)


def rule_for(target: dict) -> tuple | None:
    for rule in RULES:
        pattern, unit_head = rule[0], rule[1]
        if re.search(pattern, nfkc(target["指標"])) and base_unit(target["単位"]).startswith(unit_head):
            return rule
    return None


def reached(value: float, target: float, unit: str, how: str) -> bool:
    if how == "surplus":
        return value > 0  # 黒字化＝0を上回る（原本の「▲0.0」は −0.0 で黒字にしない）
    return value > target if "超" in unit else value >= target - 1e-9


def judge(target: dict, table: dict[str, dict[int, dict]]) -> dict:
    out = {k: target.get(k, "") for k in ("決定ID", "指標", "目標値", "単位", "目標年", "年の別")}
    rule = rule_for(target)
    if rule is None or rule[2] is None:
        return {**out, "系列ID": "", "達成の値": "", "達成の年": "", "判定": "結ばない", "理由": rule[4] if rule else DEFAULT_REASON}
    _, _, sid, how, note = rule
    year = int(target["目標年"]) if target.get("目標年") else None
    if sid == "?gdp":
        sid = GDP_CY if target.get("年の別") == "年" else GDP_FY
    s = table.get(sid, {})
    if not s:
        return {**out, "系列ID": sid, "達成の値": "", "達成の年": "", "判定": "値なし", "理由": f"{note}。系列が取り込めていない"}
    unit = target["単位"]
    goal = float(target["目標値"]) * scale(unit)
    if how == "avg10":
        first, last = year - 10, year  # 例：2013〜2022年度の平均＝2012年度から2022年度への年率
        note = f"{note}（{first + 1}〜{last}年度）"
        if first in s and last in s:
            v = ((float(s[last]["値"]) / float(s[first]["値"])) ** (1 / (last - first)) - 1) * 100
            ok = reached(v, goal, unit, how)
            return {**out, "系列ID": sid, "達成の値": round(v, 2), "達成の年": last, "判定": "達成" if ok else "未達", "理由": note}
        return {**out, "系列ID": sid, "達成の値": "", "達成の年": "", "判定": "期限前", "理由": note}
    decided = int(m.group(1)) if (m := re.search(r"(\d{4})", target["決定ID"])) else 0
    if how == "rate" and year is None:
        span = [y for y in sorted(s) if y > decided]
        if not span:
            return {**out, "系列ID": sid, "達成の値": "", "達成の年": "", "判定": "期限なし（未到達）", "理由": f"{note}。決定の翌年からの値がまだ無い"}
        v = (math.prod(1 + float(s[y]["値"]) / 100 for y in span) ** (1 / len(span)) - 1) * 100
        ok = reached(v, goal, unit, how)
        return {**out, "系列ID": sid, "達成の値": round(v, 2), "達成の年": span[-1], "判定": "期限なし（到達）" if ok else "期限なし（未到達）",
                "理由": f"{note}。{span[0]}〜{span[-1]}年の平均（決定の翌年から最新まで）"}
    years = sorted(s)
    # 届いた年は決定の年から探す（決定より前に届いていた年を「到達」としない。率の平均は決定の翌年から）
    hit = next((y for y in years if y >= decided and reached(float(s[y]["値"]), goal, unit, how)), None)

    def shown(y):
        return round(float(s[y]["値"]) / (CHO if how == "surplus" else scale(unit)), 4)

    if year is not None and year in s:
        ok = reached(float(s[year]["値"]), goal, unit, how)
        return {**out, "系列ID": sid, "達成の値": shown(year), "達成の年": year, "判定": "達成" if ok else "未達", "理由": note}
    latest = years[-1]
    if year is not None and ended(year, target.get("年の別", "")):
        return {**out, "系列ID": sid, "達成の値": shown(latest), "達成の年": latest, "判定": "値なし（期限後・実績待ち）",
                "理由": f"{note}。目標の年は終わったが、その年の実績がまだ系列に無い（最新の年の値）"}
    if year is not None and year < latest:
        return {**out, "系列ID": sid, "達成の値": "", "達成の年": "", "判定": "値なし", "理由": f"{note}。目標の年の値が系列に無い"}
    if hit is not None:  # 期限前（か期限なし）で、もう届いた年がある：最初に届いた年
        return {**out, "系列ID": sid, "達成の値": shown(hit), "達成の年": hit,
                "判定": ("期限前" if year else "期限なし") + "（到達）", "理由": f"{note}。決定の年から見て最初に届いた年"}
    return {**out, "系列ID": sid, "達成の値": shown(latest), "達成の年": latest,
            "判定": ("期限前" if year else "期限なし") + "（未到達）", "理由": f"{note}。最新の年の値"}


def ended(year: int, kind: str, today: date | None = None) -> bool:
    """目標の年（年度なら翌年3月末、年なら12月末）が終わったか。"""
    today = today or date.today()
    return today > (date(year + 1, 3, 31) if kind == "年度" else date(year, 12, 31))


def link() -> dict:
    table = series_table()
    derived = derive(table)
    write_csv(OUT / "derived.csv", SERIES_HEADER, derived)
    for r in derived:
        table.setdefault(r["系列ID"], {})[int(r["年"])] = {k: str(v) for k, v in r.items()}
    results = [judge(t, table) for t in read_csv(TARGETS)]
    write_csv(OUT / "target-links.csv", LINK_HEADER, results)
    counts: dict[str, int] = {}
    for r in results:
        key = r["判定"].split("（")[0]
        counts[key] = counts.get(key, 0) + 1
    return {"目標": len(results), **counts}


# ---------- 入口 ----------

FETCH_PHASES = ["sna", "pb", "wage", "visitors", "spending", "lodging", "fdi", "exports", "rd", "ghg"]
PHASES = FETCH_PHASES + ["link"]
STEPS = {"sna": fetch_sna, "pb": fetch_pb, "wage": fetch_wage, "visitors": fetch_visitors, "spending": fetch_spending, "lodging": fetch_lodging,
         "fdi": fetch_fdi, "exports": fetch_exports, "rd": fetch_rd, "ghg": fetch_ghg}
HOSTS = sorted({urlsplit(u).netloc for u in (SNA_TOP, SHISAN_PAGE, WAGE_PAGE, JNTO_PAGE, SPENDING_PAGE, IIP_XLS, EXPORTS_PAGE, RD_API, GHG_PAGE)})


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="国策の数値目標の達成状況を測る系列を *.go.jp の一次資料から取り込み、目標と結び付ける")
    parser.add_argument("phase", choices=PHASES + ["all"], help="回す段（all は順にすべて）")
    parser.add_argument("--monthly", action="store_true", help="段ごとに前回の取得から28日たっていなければ飛ばす（自動更新用）")
    parser.add_argument("--offline", action="store_true", help="通信せず .local/target-series/raw/ の原本だけで作り直す")
    parser.add_argument("--dry-run", action="store_true", help="通信もファイルの書き込みもせず、やることだけ出す")
    args = parser.parse_args(argv)
    phases = PHASES if args.phase == "all" else [args.phase]
    if args.monthly and not args.offline:
        phases = [p for p in phases if p not in FETCH_PHASES or due(p)]
    if "link" not in phases:
        phases.append("link")  # 取り直した系列で結び付けをいつも作り直す
    if args.dry_run:
        print(f"回す段：{', '.join(phases)}（出力 {OUT.relative_to(ROOT)}、原本 {RAW.relative_to(ROOT)}、通信先 {', '.join(HOSTS)}）")
        return 0
    OUT.mkdir(parents=True, exist_ok=True)
    fetcher = Fetcher(offline=args.offline)
    failed, fetched = False, {}
    for phase in phases:
        try:
            result = link() if phase == "link" else STEPS[phase](fetcher)
            print(f"{phase}：{result}")
            ok = True
        except Exception as error:  # 1つの取得元が落ちても残りは回す
            failed, ok = True, False
            print(f"{phase}：失敗 {type(error).__name__}: {error}")
        if phase in FETCH_PHASES:
            fetched[phase] = ok
    if not args.offline:
        fetcher.save_manifest()
        record_fetch(fetched, stamp=LAST_FETCH)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
